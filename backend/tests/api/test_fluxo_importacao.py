"""Fluxo ponta a ponta: cadastro → cálculo → recálculo → provisão → Excel → auditoria."""

import io
import re
from decimal import Decimal
from itertools import count

import pytest
from openpyxl import load_workbook

_seq = count(100)


@pytest.fixture
def payload(refs):
    n = next(_seq)
    return {
        "tipo_declaracao": "DI",
        "numero_declaracao": f"26/{n:07d}-0",
        "data_declaracao": "2026-10-16",
        "tipo_importacao_id": refs["tipos"]["USO_CONSUMO"],
        "modal": "MARITIMO",
        "valor_invoice": "10000.00",
        "moeda_invoice": "USD",
        "valor_frete": "2000.00",
        "moeda_frete": "USD",
        "qtd_itens": 10,
        "qtd_containers": 2,
        "taxa_usd": "5.000000",
        "valor_siscomex": "154.23",
        "peso_bruto_kg": "12000",
        "peso_liquido_kg": "11000",
        "transportadora_local_id": refs["transp"]["Rodo Norte Logística"],
        "transportadora_secundaria_id": refs["transp"]["Via Rio Negro Cargas"],
        "despachante_id": refs["despachante"],
        "terminal_id": refs["term"]["Chibatão"],
    }


def _criar(client, headers, payload):
    r = client.post("/api/v1/importacoes", json=payload, headers=headers)
    assert r.status_code == 201, r.text
    return r.json()


def test_login_invalido(client):
    r = client.post("/api/v1/auth/login", json={"email": "pcm@forecast.local", "senha": "errada"})
    assert r.status_code == 401


def test_sem_token(client):
    assert client.get("/api/v1/importacoes").status_code == 401


def test_cadastro_calculo_e_valores(client, pcm, payload):
    imp = _criar(client, pcm, payload)
    assert imp["status"] == "RASCUNHO"
    assert imp["numero_processo"].startswith("IMP-")

    previa = client.post(f"/api/v1/importacoes/{imp['id']}/calculos/previa", headers=pcm).json()
    assert previa["exige_motivo"] is False

    r = client.post(f"/api/v1/importacoes/{imp['id']}/calculos", json={}, headers=pcm)
    assert r.status_code == 201, r.text
    calc = r.json()
    itens = {i["componente_codigo"]: i for i in calc["itens"]}
    # Mesmos valores do caso de referência do motor, com ICMS de 20% (alíquota de exemplo)
    assert itens["CIF"]["valor"] == "60154.23"
    assert itens["II"]["valor"] == "8421.59"
    assert itens["DESPACHANTE"]["valor"] == "900.00"  # Uso e Consumo
    assert itens["TRANSP_SEC"]["valor"] == "2006.02"
    assert itens["TCIF"]["valor"] == "695.00"
    assert itens["TCIF"]["competencia"] == "2026-11-01"  # DI no dia 16
    assert itens["II"]["competencia"] == "2026-10-01"
    assert itens["CIF"]["compoe_provisao"] is False
    assert "origens" in itens["DESPACHANTE"]["memoria"]
    assert len(calc["hash_integridade"]) == 64

    detalhe = client.get(f"/api/v1/importacoes/{imp['id']}", headers=pcm).json()
    assert detalhe["status"] == "CALCULADO"
    assert detalhe["calculo_vigente"]["numero_versao"] == 1


def test_recalculo_versiona_e_exige_motivo(client, pcm, payload):
    imp = _criar(client, pcm, payload)
    client.post(f"/api/v1/importacoes/{imp['id']}/calculos", json={}, headers=pcm)
    detalhe = client.get(f"/api/v1/importacoes/{imp['id']}", headers=pcm).json()
    total_v1 = Decimal(detalhe["calculo_vigente"]["total_provisionado"])

    alterado = {**payload, "taxa_usd": "5.500000", "versao": detalhe["versao"]}
    r = client.put(f"/api/v1/importacoes/{imp['id']}", json=alterado, headers=pcm)
    assert r.status_code == 200, r.text
    assert r.json()["calculo_desatualizado"] is True

    sem_motivo = client.post(f"/api/v1/importacoes/{imp['id']}/calculos", json={}, headers=pcm)
    assert sem_motivo.status_code == 422

    previa = client.post(f"/api/v1/importacoes/{imp['id']}/calculos/previa", headers=pcm).json()
    assert previa["exige_motivo"] is True
    assert Decimal(previa["diferencas"]["CIF"]["diferenca"]) == Decimal("6000.00")

    r = client.post(
        f"/api/v1/importacoes/{imp['id']}/calculos",
        json={"motivo": "Atualização PTAX"},
        headers=pcm,
    )
    assert r.status_code == 201
    assert Decimal(r.json()["total_provisionado"]) > total_v1

    hist = client.get(f"/api/v1/importacoes/{imp['id']}/calculos", headers=pcm).json()
    assert [h["numero_versao"] for h in hist] == [2, 1]
    assert [h["vigente"] for h in hist] == [True, False]


def test_conflito_de_edicao(client, pcm, payload):
    imp = _criar(client, pcm, payload)
    desatualizado = {
        **payload,
        "observacoes": "x",
        "versao": imp["versao"] - 1 if imp["versao"] > 1 else 999,
    }
    r = client.put(f"/api/v1/importacoes/{imp['id']}", json=desatualizado, headers=pcm)
    assert r.status_code == 409


def test_validacoes_de_entrada(client, pcm, payload):
    sem_taxa = {**payload, "moeda_invoice": "CNY", "taxa_cny": None}
    r = client.post("/api/v1/importacoes", json=sem_taxa, headers=pcm)
    assert r.status_code == 422
    assert "taxa CNY" in r.text

    di_ruim = {**payload, "numero_declaracao": "123"}
    assert client.post("/api/v1/importacoes", json=di_ruim, headers=pcm).status_code == 422

    sem_data = {**payload, "data_declaracao": None}
    assert client.post("/api/v1/importacoes", json=sem_data, headers=pcm).status_code == 422


def test_di_duplicada(client, pcm, payload):
    _criar(client, pcm, payload)
    r = client.post("/api/v1/importacoes", json=payload, headers=pcm)
    assert r.status_code == 409
    assert "IMP-" in r.json()["detail"]


def test_financeiro_nao_cadastra_nem_calcula(client, financeiro, pcm, payload):
    assert client.post("/api/v1/importacoes", json=payload, headers=financeiro).status_code == 403
    imp = _criar(client, pcm, payload)
    r = client.post(f"/api/v1/importacoes/{imp['id']}/calculos", json={}, headers=financeiro)
    assert r.status_code == 403
    assert client.get("/api/v1/importacoes", headers=financeiro).status_code == 200


def test_provisao_cancelamento_e_excel(client, pcm, financeiro, payload):
    payload = {**payload, "data_declaracao": "2027-03-10"}  # mês isolado
    imp = _criar(client, pcm, payload)
    client.post(f"/api/v1/importacoes/{imp['id']}/calculos", json={}, headers=pcm)

    params = {"de": "2027-03-01", "ate": "2027-03-01"}
    mensal = client.get("/api/v1/provisoes/mensal", params=params, headers=financeiro).json()
    codigos = {ln["chave"] for ln in mensal["linhas"]}
    assert "CIF" not in codigos and "II" in codigos
    assert re.fullmatch(r"\d+\.\d{2}", mensal["total"])  # string decimal, sem float
    total = Decimal(mensal["total"])
    assert total > 0

    por_processo = client.get(
        "/api/v1/provisoes/por-processo", params=params, headers=financeiro
    ).json()
    assert por_processo["linhas"][0]["rotulo"] == imp["numero_processo"]
    assert Decimal(por_processo["total"]) == total

    por_tipo = client.get("/api/v1/provisoes/por-tipo", params=params, headers=financeiro).json()
    assert por_tipo["linhas"][0]["chave"] == "USO_CONSUMO"

    r = client.get("/api/v1/provisoes/mensal/excel", params=params, headers=financeiro)
    assert r.status_code == 200
    planilha = load_workbook(io.BytesIO(r.content)).active
    ultima = planilha.max_row
    assert planilha.cell(ultima, 1).value == "Total"
    # Excel armazena números em ponto flutuante; comparamos em centavos
    valor_excel = Decimal(str(planilha.cell(ultima, planilha.max_column).value))
    assert valor_excel.quantize(Decimal("0.01")) == total

    r = client.post(
        f"/api/v1/importacoes/{imp['id']}/cancelar",
        json={"motivo": "Pedido cancelado"},
        headers=pcm,
    )
    assert r.json()["status"] == "CANCELADO"
    depois = client.get("/api/v1/provisoes/mensal", params=params, headers=financeiro).json()
    assert Decimal(depois["total"]) == 0


def test_auditoria_registra_usuario_e_campos(client, pcm, admin, payload):
    imp = _criar(client, pcm, payload)
    alterado = {**payload, "observacoes": "Revisado", "versao": imp["versao"]}
    client.put(f"/api/v1/importacoes/{imp['id']}", json=alterado, headers=pcm)
    r = client.get(
        "/api/v1/auditoria",
        params={"tabela": "importacao", "registro_id": imp["id"]},
        headers=admin,
    ).json()
    acoes = [i["acao"] for i in r["itens"]]
    assert acoes == ["UPDATE", "INSERT"]
    assert r["itens"][0]["usuario_nome"] == "Analista PCM"
    assert r["itens"][0]["campos_alterados"] == ["observacoes"]


def test_dashboard(client, financeiro):
    r = client.get("/api/v1/dashboard", headers=financeiro)
    assert r.status_code == 200
    corpo = r.json()
    assert {"kpis", "serie", "importacoes_por_mes", "por_tipo", "composicao"} <= corpo.keys()
    assert len(corpo["serie"]) == 12
